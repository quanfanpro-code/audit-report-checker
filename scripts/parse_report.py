"""本地提取原始资料并保留缺失清单；不自动调用模型或云端。"""
import argparse
import re
import shutil
import subprocess
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree
from 复核契约 import save_json,save_text,fingerprint,backup_existing

W="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SUPPORTED={".docx",".pdf",".xlsx",".xlsm",".png",".jpg",".jpeg",".tif",".tiff",".bmp",".doc"}

def _word_text(element,owner_table=None):
    parts=[]
    for node in element.iter():
        if owner_table is not None:
            parent=node.getparent()
            while parent is not None and parent.tag!=W+"tbl":parent=parent.getparent()
            if parent is not owner_table:continue
        if node.tag==W+"p" and parts and parts[-1]!="\n":parts.append("\n")
        elif node.tag==W+"t":parts.append(node.text or "")
        elif node.tag==W+"tab":parts.append("\t")
        elif node.tag in (W+"br",W+"cr"):parts.append("\n")
    return "".join(parts)

def render_md_table(headers,rows):
    def escape(x):return str(x if x is not None else "[缺失]").replace("|","\\|").replace("\n"," ")
    lines=[]
    if headers:lines.append("| "+" | ".join(map(escape,headers))+" |")
    for i,row in enumerate(rows):lines.append(str(i)+": | "+" | ".join(map(escape,row))+" |")
    return "\n".join(lines)

def _sdt_aware_extract_tbl(element):
    rows=[]
    def nearest(node,tag):
        parent=node.getparent()
        while parent is not None and parent.tag!=tag:parent=parent.getparent()
        return parent
    for tr in element.iter(W+"tr"):
        if nearest(tr,W+"tbl") is not element:continue
        cells=[]
        for tc in tr.iter(W+"tc"):
            if nearest(tc,W+"tr") is not tr:continue
            span=tc.find(W+"tcPr/"+W+"gridSpan")
            size=int(span.get(W+"val","1")) if span is not None else 1
            text=_word_text(tc,element)
            cells.extend([text]+[""]*(size-1))
        rows.append(cells)
    return rows

def _process_docx_text(path,rel,out_dir=None):
    from docx import Document
    doc=Document(str(path));blocks=[];tables=[];pending=[]
    heading=""
    def walk(node):
        nonlocal heading
        if node.tag==W+"p":
            text=_word_text(node)
            if text:
                style=node.find(W+"pPr/"+W+"pStyle")
                if style is not None and re.search("Heading|标题",style.get(W+"val",""),re.I):heading=text
                blocks.append({"text":text,"chapter":heading,"part":"正文"})
        elif node.tag==W+"tbl":
            rows=_sdt_aware_extract_tbl(node)
            table={"headers":[],"rows":rows,"row_model":"all_raw","source_file":rel,"chapter":heading,"page":None,"method":"python-docx","name":""}
            tables.append(table);blocks.append({"table":table,"chapter":heading,"part":"正文表格"})
            for nested in node.findall(".//"+W+"tbl"):
                parent=nested.getparent()
                while parent is not None and parent.tag!=W+"tbl":parent=parent.getparent()
                if parent is node:walk(nested)
        else:
            for child in node:walk(child)
    walk(doc.element.body)
    with zipfile.ZipFile(path) as package:
        for part in package.namelist():
            if re.fullmatch(r"word/(?:header|footer|footnotes|endnotes)\d*\.xml",part):
                for para in ElementTree.fromstring(package.read(part)).iter(W+"p"):
                    text=_word_text(para)
                    if text:blocks.append({"text":text,"part":part,"chapter":part})
            if part.startswith("word/media/"):
                pending.append({"source_file":rel,"part":part,"reason":"嵌入图片需当前agent识别用途和内容"})
                if out_dir is not None:
                    dest=Path(out_dir)/"图片"/(Path(rel).stem+"_"+Path(part).name)
                    dest.parent.mkdir(parents=True,exist_ok=True)
                    if dest.exists():backup_existing(dest)
                    payload=package.read(part);dest.write_bytes(payload)
                    if dest.read_bytes()!=payload:raise OSError("图片写入回读失败")
                    pending[-1]["image_path"]=str(dest.resolve())
        if b"<w:del" in package.read("word/document.xml"):
            pending.append({"source_file":rel,"reason":"存在修订删除内容，需确认版本"})
    return blocks,tables,pending

def _process_pdf_text(path,rel):
    import pdfplumber
    blocks=[];tables=[];pending=[]
    with pdfplumber.open(path) as pdf:
        for number,page in enumerate(pdf.pages,1):
            text=page.extract_text() or ""
            blocks.append({"text":text,"page":number,"part":"页面文字"})
            for obj in page.find_tables():
                table={"headers":[],"rows":obj.extract(),"row_model":"all_raw","source_file":rel,"page":number,"bbox":list(obj.bbox),"method":"pdfplumber","name":""}
                tables.append(table);blocks.append({"table":table,"page":number,"part":"页面表格"})
            if not text.strip() or page.images:
                pending.append({"source_file":rel,"page":number,"reason":"图片或无可提取文字，需当前agent按原页面核对"})
    return blocks,tables,pending

def _process_xlsx(path,rel):
    import openpyxl
    cached=openpyxl.load_workbook(path,data_only=True,read_only=True)
    formulas=openpyxl.load_workbook(path,data_only=False,read_only=True)
    layout=openpyxl.load_workbook(path,read_only=False)
    blocks=[];tables=[];pending=[]
    try:
        for sn in cached.sheetnames:
            visibility=formulas[sn].sheet_state
            if visibility != "visible":
                pending.append({"source_file":rel,"sheet":sn,"visibility":visibility,"reason":"隐藏工作表默认不属于报出资料，不读取、不参与复核；如需核对请先确认并取消隐藏"})
                continue
            hidden_rows=set();hidden_cols=set()
            sheet_layout=layout[sn]
            for idx,dim in sheet_layout.row_dimensions.items():
                if dim.hidden:hidden_rows.add(idx-1)
            for letter,dim in sheet_layout.column_dimensions.items():
                if dim.hidden:
                    from openpyxl.utils import column_index_from_string
                    start=dim.min or column_index_from_string(letter)
                    hidden_cols.update(range(start-1,(dim.max or start)))
            if hidden_rows or hidden_cols:
                pending.append({"source_file":rel,"sheet":sn,"reason":"隐藏行或隐藏列默认不属于报出资料，不读取、不参与复核；如需核对请先确认并取消隐藏"})
            rows=[];formula_cells={}
            cached_rows=iter(cached[sn].iter_rows(values_only=True))
            for r,row in enumerate(formulas[sn].iter_rows()):
                cache_row=next(cached_rows,())
                if r in hidden_rows:continue
                values=[]
                for c,cell in enumerate(row):
                    if c in hidden_cols:continue
                    value=cache_row[c] if c<len(cache_row) else None
                    if cell.data_type=="f":
                        formula_cells[f"{r},{c}"]={"formula":cell.value,"cached":value}
                        if value is None:pending.append({"source_file":rel,"sheet":sn,"row":r,"col":c,"reason":"公式缺少缓存值，不能当零"})
                    values.append(str(value) if value is not None else None)
                rows.append(values)
            table={"headers":[],"rows":rows,"row_model":"all_raw","source_file":rel,"sheet":sn,"page":None,"method":"openpyxl","name":sn,"formulas":formula_cells}
            tables.append(table);blocks.append({"table":table,"part":sn})
    finally:cached.close();formulas.close();layout.close()
    return blocks,tables,pending

def _process_mineru(path,rel,out_dir,*,allow_cloud_ocr=False,pages=None):
    if not allow_cloud_ocr:raise ValueError("未明确授权本次上传，云端OCR不会执行")
    command=shutil.which("mineru-open-api")
    if not command:raise ValueError("本机没有可用的mineru命令；未安装任何依赖")
    target=Path(path)
    if pages:
        if target.suffix.lower()!=".pdf" or not re.fullmatch(r"\d+-\d+",pages):raise ValueError("页范围必须是PDF起止页")
        import pymupdf
        begin,end=map(int,pages.split("-"))
        with pymupdf.open(target) as doc:
            if not 1<=begin<=end<=len(doc):raise ValueError("页范围越界")
            selected=pymupdf.open();selected.insert_pdf(doc,from_page=begin-1,to_page=end-1)
            target=Path(out_dir)/("授权选页_"+uuid.uuid4().hex+".pdf");selected.save(target);selected.close()
    cloud_out=Path(out_dir)/("授权OCR_"+uuid.uuid4().hex[:12])
    result=subprocess.run([command,"extract",str(target),"-o",str(cloud_out),"-f","md","--model","vlm","--language","ch"],capture_output=True,text=True,encoding="utf-8",timeout=1800)
    if result.returncode:raise ValueError("OCR调用失败："+result.stderr[:300])
    paths=sorted(cloud_out.rglob("*.md"))
    if not paths:raise ValueError("OCR没有返回正文")
    return "\n\n".join(p.read_text(encoding="utf-8-sig") for p in paths)

def main(argv=None):
    parser=argparse.ArgumentParser(description="只读本地提取及缺失登记")
    parser.add_argument("input");parser.add_argument("-o","--output",required=True)
    parser.add_argument("--allow-cloud-ocr",action="store_true",help="仅在用户明确授权本次单文件上传后使用")
    parser.add_argument("--ocr-pages",help="授权PDF页范围，先本地拆页再上传")
    args=parser.parse_args(argv)
    source=Path(args.input).resolve();out=Path(args.output).resolve()
    if not source.exists():parser.error("输入不存在")
    if out==source or (source.is_file() and out==source.parent):parser.error("输出必须使用独立目录")
    if args.allow_cloud_ocr and source.is_dir():parser.error("云端OCR授权限于本次明确的单文件")
    if args.ocr_pages and not args.allow_cloud_ocr:parser.error("选择OCR页范围也需要明确上传授权")
    files=sorted((p for p in source.iterdir() if p.is_file() and not p.name.startswith("~$")),key=lambda p:p.name) if source.is_dir() else [source]
    if not files:parser.error("没有支持的来源文件")
    root=source if source.is_dir() else source.parent
    out.mkdir(parents=True,exist_ok=True)
    data={"schema_version":2,"source_root":str(root),"sources":[],"files":[],"tables":[],"segments":[],"pending":[],"company":"","period":""}
    parts=[]
    for file in files:
        rel=str(file.relative_to(root));sha=fingerprint(file)
        data["sources"].append(rel);info={"name":rel,"path":str(file),"sha256":sha,"status":"已提取"};data["files"].append(info)
        try:
            if file.suffix.lower()==".docx":blocks,tables,pending=_process_docx_text(file,rel,out)
            elif file.suffix.lower()==".pdf":blocks,tables,pending=_process_pdf_text(file,rel)
            elif file.suffix.lower() in (".xlsx",".xlsm"):blocks,tables,pending=_process_xlsx(file,rel)
            else:blocks=[];tables=[];pending=[{"source_file":rel,"reason":"需当前agent读取图片或提供可读文档版本" if file.suffix.lower() in SUPPORTED else "不支持的文件格式，请转换为可读版本："+file.suffix,"image_path":str(file)}]
            if args.allow_cloud_ocr:
                try:blocks.append({"text":_process_mineru(file,rel,out,allow_cloud_ocr=True,pages=args.ocr_pages),"part":"授权OCR，尚需回原件核对"})
                except Exception as exc:pending.append({"source_file":rel,"reason":"OCR补充失败，本地提取保留："+str(exc)})
            for table in tables:
                table["id"]=len(data["tables"]);data["tables"].append(table)
            for block in blocks:
                text=render_md_table([],block["table"]["rows"]) if "table" in block else block.get("text","")
                marker=f'<!-- SOURCE file="{rel}"'+(f' page={block["page"]}' if block.get("page") else "")+' -->'
                parts.append(marker+"\n"+text)
                segment={"id":len(data["segments"]),"source_file":rel,"part":block.get("part",""),"page":block.get("page"),"chapter":block.get("chapter",""),"text":text}
                if "table" in block:segment["table_id"]=block["table"]["id"]
                data["segments"].append(segment)
            data["pending"].extend(pending)
            if pending:info["status"]="部分待核实"
        except Exception as exc:
            info["status"]="提取失败";data["pending"].append({"source_file":rel,"reason":type(exc).__name__+": "+str(exc)})
        if fingerprint(file)!=sha:raise ValueError("读取期间源文件发生变化："+rel)
    save_text(out/"report.md","\n\n".join(parts)+"\n")
    save_json(out/"extracted_tables.json",data)
    save_json(out/"files_index.json",{i["name"]:{"path":i["path"],"loader":Path(i["name"]).suffix[1:],"sha256":i["sha256"]} for i in data["files"]})
    print("本地提取完成；待核实项目："+str(len(data["pending"]))+"；尚未完成审计复核")
    return 0

if __name__=="__main__":raise SystemExit(main())
