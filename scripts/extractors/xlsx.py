"""读取Excel缓存值并核查公式，不自动改动或重算源文件。"""
from pathlib import Path
import openpyxl

def load_xlsx_sheets(path):
    path=Path(path)
    cached=openpyxl.load_workbook(path,data_only=True)
    formulas=openpyxl.load_workbook(path,data_only=False)
    try:
        for sn in formulas.sheetnames:
            for row in formulas[sn].iter_rows():
                for cell in row:
                    if cell.data_type=="f" and cached[sn][cell.coordinate].value is None:
                        raise ValueError("公式缺少缓存值："+sn+"!"+cell.coordinate)
    finally:formulas.close()
    return {sn:cached[sn] for sn in cached.sheetnames}

