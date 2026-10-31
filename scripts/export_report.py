"""导出人类可读的问题报告，并在隐藏工作表中保留技术留痕。"""
import argparse
import math
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from 复核契约 import load_json, validate_report, save_text, backup_existing, verify_sources, guard_outputs

SHEETS = ["问题概览", "问题明细", "技术留痕"]
SEVERITY_ORDER = {"致命": 0, "重要": 1, "细节": 2}
STATUS_SEVERITY = {"确认错误": "致命", "待核实": "重要", "未执行": "重要"}
COLORS = {"致命": "F4CCCC", "重要": "FCE5CD", "细节": "E8F0FE"}
SCOPE_NAMES = {"quick": "快速复核", "standard": "标准复核", "deep": "全面复核"}


def load_results(path):
    data = load_json(path)
    verify_sources(data)
    data.setdefault("meta", {})
    data["meta"].setdefault("stage", "draft")
    data.setdefault("results", [])
    data.setdefault("coverage", {"expected_ids": []})
    validate_report(data)
    return data


def field(item, name, default=""):
    value = item.get(name, default)
    return default if value is None else str(value)


def _basis(item):
    return field(item.get("review", {}), "basis") or field(item, "description")


def _clean(value):
    text = str(value or "")
    text = re.sub(r"(?<![A-Za-z0-9_-])[a-z][a-z0-9]*(?:-[a-z0-9]+)+(?![A-Za-z0-9_-])", "", text)
    replacements = {
        "relations": "核对关系",
        "计算草稿": "验算结果",
        "本地验算": "验算",
        "浮点残留": "计算尾差",
        "终审": "复核",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"(?<=[：；])/", "", text)
    return re.sub(r"\s+", " ", text).strip(" ；;，,/")


def _severity(item):
    declared = item.get("review", {}).get("severity") or item.get("severity")
    return declared if declared in SEVERITY_ORDER else STATUS_SEVERITY.get(item.get("status"), "细节")


def _source_tables(data):
    directory = Path(field(data.get("meta", {}), "input_dir"))
    path = directory / "extracted_tables.json"
    if not directory.is_dir() or not path.is_file():
        return {}
    return {table.get("id"): table for table in load_json(path).get("tables", [])}


def _table(item, tables, table_id=None):
    key = item.get("table_id") if table_id is None else table_id
    return tables.get(key, tables.get(str(key), {}))


def _location(item, tables=None):
    tables = tables or {}
    source = Path(field(item, "source_file")).name
    table = _table(item, tables)
    parts = [source] if source else []
    for key in ("sheet", "name", "chapter"):
        value = table.get(key) or item.get(key)
        if value and str(value) not in parts:
            parts.append(str(value))
            break
    if item.get("page") is not None:
        parts.append("第" + str(item["page"]) + "页")
    if item.get("row_range") is not None:
        start, end = item["row_range"]
        parts.append("第" + str(start + 1) + ("–" + str(end) if end > start + 1 else "") + "行")
    elif item.get("source_span") is not None and item.get("table_id") is None:
        parts.append("对应文字片段")
    if not parts and item.get("table_id") is None:
        return "跨表核对汇总"
    if not table and item.get("table_id") is not None and not any(key in item for key in ("sheet", "name", "chapter")):
        parts.insert(1, "文档内第" + str(item["table_id"] + 1) + "个表格")
    return " · ".join(parts)


def _rule_group(row):
    name = re.sub(r"\([^)]*(?:预期|见复核)[^)]*\)", "", field(row, "rule_name"))
    return _clean(name.replace("本期", "").replace("上期", "")).strip(" ：:，,")


def _primary_rule(item):
    rows = _mismatches(item)
    if not rows:
        return ""
    row = next((row for row in rows if "预期存在" not in field(row, "rule_name")), rows[0])
    return _rule_group(row)


def _title(item):
    if item.get("issue_title"):
        return field(item, "issue_title")
    group = item.get("_rule_group", "")
    if group:
        group = group.replace("CFS", "现金流量表").replace("BS", "资产负债表")
        return group.replace("=", "与") + "不一致" if "=" in group else group
    basis = _clean(_basis(item))
    if basis:
        parts = [part.strip() for part in re.split(r"[。；]", basis) if part.strip()]
        selected = next((part for part in reversed(parts) if "确认错误" in part), parts[0])
        return selected if len(selected) <= 90 else selected[:89] + "…"
    name = _clean(field(item, "rule_name", "复核发现的问题"))
    return re.sub(r"^(?:table|text|cross_relations)-?", "", name) or "复核发现的问题"


def _mismatches(item):
    rows = [c for c in item.get("calculations", []) if c.get("calculation", {}).get("match") is False]
    if not rows and item.get("expected") is not None and item.get("actual") is not None and item.get("status") in STATUS_SEVERITY:
        rows = [item]
    unique = []
    seen = set()
    for row in rows:
        key = (field(row, "rule_name"), field(row, "expected"), field(row, "actual"), field(row, "difference"))
        if key not in seen:
            seen.add(key)
            unique.append(row)
    group = item.get("_rule_group")
    return [row for row in unique if _rule_group(row) == group] if group else unique


def _amount(value):
    if value in (None, ""):
        return "—"
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            return str(value)
        return format(number, ",.2f")
    except (InvalidOperation, ValueError):
        return str(value)


def _unit(item, tables):
    target = item.get("target", {})
    if target.get("unit"):
        return target["unit"]
    return _table(item, tables, target.get("table_id")).get("unit") or item.get("unit") or ""


def _comparison(item, tables):
    rows = _mismatches(item)
    lines = []
    for row in rows[:5]:
        unit = _unit(row, tables)
        lines.append(
            (_clean(field(row, "rule_name")) or "核对关系") + "：应为" + _amount(row.get("expected")) + unit
            + "，实际" + _amount(row.get("actual")) + unit + "，差异" + _amount(row.get("difference")) + unit
        )
    if len(rows) > 5:
        lines.append("另有" + str(len(rows) - 5) + "条差异，见技术留痕。")
    return "\n".join(lines) or "—"


def _human_context(value):
    text = str(value or "").strip()
    markers = ('"row_range"', '"rows"', '"kind"', "CHECK_ID", "source_fingerprint", '"span"')
    return bool(text) and not any(marker in text for marker in markers) and not text.startswith(("{", "["))


def _row_text(table, row):
    rows = table.get("rows", [])
    if type(row) is not int or not 0 <= row < len(rows):
        return ""
    return " | ".join(str(value) for value in rows[row] if value not in (None, ""))


def _original(item, tables):
    if item.get("original_text"):
        return field(item, "original_text")
    if _human_context(item.get("context")):
        return field(item, "context")
    rows = []
    for calculation in _mismatches(item):
        refs = [*calculation.get("operands", []), calculation.get("target", {})]
        for ref in refs:
            text = _row_text(_table(item, tables, ref.get("table_id")), ref.get("row"))
            if text and text not in rows:
                rows.append(text)
    if not rows and item.get("row_range") is not None:
        start, end = item["row_range"]
        table = _table(item, tables)
        rows = [text for text in (_row_text(table, row) for row in range(start, min(end, start + 2))) if text]
    if rows:
        return "\n".join(rows[:3])
    rows = _mismatches(item)
    if rows:
        first = rows[0]
        return "报告列示：" + _amount(first.get("actual")) + _unit(first, tables) + "；核对基准：" + _amount(first.get("expected")) + _unit(first, tables)
    return "—"


def _suggestion(item):
    if item.get("fix_suggestion"):
        return field(item, "fix_suggestion")
    basis = _basis(item)
    actions = re.findall(r"(?:需|请)[^。；]*", basis)
    if actions:
        return "；".join(actions)
    return {
        "确认错误": "核对并更正对应明细或合计数，使报告与核对基准一致；修改后重新复核。",
        "待核实": "由报告编制人核实该处内容，确认后修正或补充依据。",
        "未执行": "补充完成该项检查所需的资料或核对，未落实前不作为已完成。",
    }.get(item.get("status"), "结合复核依据确认是否需要处理。")


def _card(item, tables):
    basis = _clean(_basis(item))
    if item.get("_group_count", 1) > 1:
        basis += "；本项归并共" + str(item["_group_count"]) + "条同类记录"
    return {
        "severity": _severity(item),
        "title": _title(item),
        "location": _location(item, tables),
        "original": _original(item, tables),
        "comparison": _comparison(item, tables),
        "basis": basis or "—",
        "suggestion": _suggestion(item),
    }


def _cards(data, tables):
    grouped = {}
    cards = []
    for item in data["results"]:
        if item.get("status") not in STATUS_SEVERITY or item.get("use_declared_relations"):
            continue
        basis = _basis(item)
        if "无独立验算对象" in basis or "不做独立验算" in basis:
            continue
        item = dict(item)
        if item.get("status") == "确认错误":
            item["_rule_group"] = _primary_rule(item)
            key = ("确认错误", item["_rule_group"] or _title(item))
        elif item.get("status") == "未执行":
            key = ("未执行", Path(field(item, "source_file")).name, basis, _severity(item))
        else:
            key = None
        if key is not None:
            if key in grouped:
                grouped[key]["_group_count"] += 1
                continue
            item["_group_count"] = 1
            grouped[key] = item
        cards.append(item)
    cards.sort(key=lambda item: (SEVERITY_ORDER[_severity(item)], _location(item, tables)))
    return [_card(item, tables) for item in cards]


def _counts(cards):
    return {severity: sum(card["severity"] == severity for card in cards) for severity in SEVERITY_ORDER}


def _conclusion(data, counts):
    if counts["致命"]:
        return "当前不建议直接出报：先处理" + str(counts["致命"]) + "项致命问题。"
    if counts["重要"]:
        return "确认" + str(counts["重要"]) + "项重要事项后再定稿。"
    if data.get("meta", {}).get("stage") != "final":
        return "复核尚未完成，当前结果仅供阶段检查。"
    return "未发现需优先处理的问题，可按既定范围完成后续出报复核。"


def _cell(ws, row, col, value, bold=False, fill=None, size=11):
    cell = ws.cell(row, col, str(value) if value is not None else "")
    cell.data_type = "s"
    cell.font = Font(name="微软雅黑", size=size, bold=bold, color="FFFFFF" if fill == "4472C4" else "222222")
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _row_height(ws, row, widths):
    lines = 1
    for col, width in widths.items():
        text = str(ws.cell(row, col).value or "")
        wrapped = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line) * 1.15 / width)) for line in text.split("\n"))
        lines = max(lines, wrapped)
    ws.row_dimensions[row].height = min(409, max(24, lines * 19 + 8))


def _write_overview(wb, data, cards, counts):
    ws = wb.active
    ws.title = SHEETS[0]
    company = field(data.get("meta", {}), "company", "审计报告")
    ws.merge_cells("A1:D1")
    _cell(ws, 1, 1, company + "复核结果", True, "4472C4", 16)
    ws.merge_cells("A2:D2")
    _cell(ws, 2, 1, _conclusion(data, counts), True, COLORS["致命"] if counts["致命"] else COLORS["重要"], 12)
    for col, label in enumerate(("致命", "重要", "细节", "已排除误报"), 1):
        _cell(ws, 4, col, label, True, "D9E2F3")
        value = counts.get(label, 0) if label != "已排除误报" else sum(item.get("status") == "已排除误报" for item in data["results"])
        _cell(ws, 5, col, value, True, COLORS.get(label), 15)
    ws.merge_cells("A7:D7")
    _cell(ws, 7, 1, "最需要先处理的问题", True, "D9E2F3")
    priorities = [card for card in cards if card["severity"] == "致命"] or [card for card in cards if card["severity"] == "重要"]
    row = 8
    for card in priorities[:5]:
        _cell(ws, row, 1, card["severity"], True, COLORS[card["severity"]])
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        _cell(ws, row, 2, card["title"] + "（" + card["location"] + "）")
        _row_height(ws, row, {2: 100})
        row += 1
    if not priorities:
        ws.merge_cells("A8:D8")
        _cell(ws, 8, 1, "当前没有需要优先展示的问题。")
        row = 9
    row += 1
    scope = SCOPE_NAMES.get(field(data.get("meta", {}), "scope"), field(data.get("meta", {}), "scope", "按本次选择范围"))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    _cell(ws, row, 1, "复核范围：" + str(len(data.get("meta", {}).get("source_files", []))) + "份文件｜" + scope + "｜原始记录" + str(len(data["results"])) + "条")
    for col, width in zip("ABCD", (18, 32, 32, 24)):
        ws.column_dimensions[col].width = width
    for number in (1, 2, 4, 5, row):
        _row_height(ws, number, {1: 100})
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def _write_details(wb, cards):
    ws = wb.create_sheet(SHEETS[1])
    ws.merge_cells("A1:B1")
    _cell(ws, 1, 1, "问题明细｜按致命、重要、细节排序；每项均可直接定位和修改", True, "4472C4", 12)
    row = 3
    for card in cards:
        values = (
            ("问题", card["severity"] + "｜" + card["title"]),
            ("错在哪", card["location"]),
            ("原文", card["original"]),
            ("核对结果", card["comparison"]),
            ("凭什么说错", card["basis"]),
            ("怎么改", card["suggestion"]),
        )
        for label, value in values:
            _cell(ws, row, 1, label, True, COLORS[card["severity"]] if label == "问题" else "D9E2F3")
            _cell(ws, row, 2, value, label == "问题", COLORS[card["severity"]] if label == "问题" else None)
            _row_height(ws, row, {2: 105})
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"


def _write_technical(wb, data, results_path):
    ws = wb.create_sheet(SHEETS[2])
    headings = ["对象ID", "状态", "检查类型", "内部规则名", "完整文件路径", "表格ID", "行号区间", "字符区间", "来源指纹", "结果文件路径", "原始复核依据", "原始上下文"]
    for col, value in enumerate(headings, 1):
        _cell(ws, 1, col, value, True, "4472C4")
    for row, item in enumerate(data["results"], 2):
        values = [
            item["check_id"], item["status"], field(item, "check_type"), field(item, "rule_name"), field(item, "source_file"),
            field(item, "table_id"), field(item, "row_range"), field(item, "source_span"), field(item, "source_fingerprint"),
            str(Path(results_path).resolve()), _basis(item), field(item, "context"),
        ]
        for col, value in enumerate(values, 1):
            _cell(ws, row, col, value)
        _row_height(ws, row, {11: 70, 12: 45})
    for col, width in zip("ABCDEFGHIJKL", (28, 14, 16, 36, 65, 12, 18, 18, 30, 65, 75, 50)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_state = "hidden"


def _protect(data, out_path, results_path=None):
    output = Path(out_path).resolve()
    if data.get("meta", {}).get("input_dir"):
        guard_outputs(data["meta"]["input_dir"], output)
    protected = {Path(results_path).resolve()} if results_path is not None else set()
    protected.update(Path(f["path"]).resolve() for f in data.get("meta", {}).get("source_files", []) if f.get("path"))
    for item in data.get("results", []):
        if item.get("source_file"):
            protected.add(Path(item["source_file"]).resolve())
        for ref in item.get("source_refs", []):
            if isinstance(ref, dict) and ref.get("source_file"):
                protected.add(Path(ref["source_file"]).resolve())
    if output in protected:
        raise ValueError("输出不能覆盖原始来源或结果输入")
    return output


def write_excel(data, results_path, out_path):
    verify_sources(data)
    validate_report(data)
    out = _protect(data, out_path, results_path)
    tables = _source_tables(data)
    cards = _cards(data, tables)
    counts = _counts(cards)
    wb = Workbook()
    _write_overview(wb, data, cards, counts)
    _write_details(wb, cards)
    _write_technical(wb, data, results_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        backup_existing(out)
    wb.save(out)
    wb.close()


def write_markdown(data, out_path):
    verify_sources(data)
    validate_report(data)
    out = _protect(data, out_path)
    tables = _source_tables(data)
    cards = _cards(data, tables)
    counts = _counts(cards)
    company = field(data.get("meta", {}), "company", "审计报告")
    false_count = sum(item.get("status") == "已排除误报" for item in data["results"])
    lines = [
        "# " + company + "复核结果", "", "> " + _conclusion(data, counts), "",
        "## 问题概览", "",
        "- 致命：" + str(counts["致命"]) + "项",
        "- 重要：" + str(counts["重要"]) + "项",
        "- 细节：" + str(counts["细节"]) + "项",
        "- 已排除误报：" + str(false_count) + "条，不列入待办", "",
    ]
    priorities = [card for card in cards if card["severity"] == "致命"] or [card for card in cards if card["severity"] == "重要"]
    if priorities:
        lines.append("最需要先处理的问题：")
        lines.extend("- **" + card["severity"] + "｜" + card["title"] + "**：" + card["location"] for card in priorities[:5])
    lines.extend(["", "## 问题明细", ""])
    for card in cards:
        lines.extend([
            "### " + card["severity"] + "｜" + card["title"], "",
            "- **错在哪**：" + card["location"],
            "- **原文**：" + card["original"].replace("\n", "；"),
            "- **核对结果**：" + card["comparison"].replace("\n", "；"),
            "- **凭什么说错**：" + card["basis"],
            "- **怎么改**：" + card["suggestion"], "",
        ])
    save_text(out, "\n".join(lines).rstrip() + "\n")


def derive_output_paths(results_path, meta, output_arg):
    stem = Path(str(meta.get("source") or "审计报告")).stem
    for ch in '/\\:*?"<>|':
        stem = stem.replace(ch, "_")
    if output_arg and Path(output_arg).suffix.lower() == ".xlsx":
        path = Path(output_arg).resolve()
        return path, path.with_name(path.stem + "_复核报告.md")
    directory = Path(output_arg).resolve() if output_arg else Path(results_path).parent
    return directory / (stem + "_检查报告.xlsx"), directory / (stem + "_复核报告.md")


def main(argv=None):
    parser = argparse.ArgumentParser(description="导出问题概览、问题明细、Markdown及默认隐藏的技术留痕")
    parser.add_argument("results")
    parser.add_argument("-o", "--output")
    args = parser.parse_args(argv)
    try:
        path = Path(args.results).resolve()
        data = load_results(path)
        excel, markdown = derive_output_paths(path, data["meta"], args.output)
        write_excel(data, path, excel)
        write_markdown(data, markdown)
        print("已导出：" + str(excel) + "；" + str(markdown))
        return 0
    except (ValueError, KeyError, TypeError, OSError) as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    raise SystemExit(main())
