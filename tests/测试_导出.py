from pathlib import Path
import pytest
from openpyxl import load_workbook
import export_report as export


def test_空结果不能报告通过(验证目录):
    path = 验证目录 / "空结果.md"
    export.write_markdown({"meta": {"stage": "draft"}, "coverage": {"expected_ids": ["a"]}, "results": []}, path)
    text = path.read_text(encoding="utf-8-sig")
    assert "检查通过" not in text and "未完成" in text


def record(cid, status):
    return {"check_id": cid, "status": status, "rule_name": "检查" + cid, "description": "核对原文", "check_type": "表注", "calculation": {"match": status == "通过"}, "source_file": "C:/样例/报告.pdf", "page": 2, "context": "原文摘录", "expected": "100.00", "actual": "99.00", "difference": "1.00", "review": {"agent": "当前agent", "check_id": cid, "source_fingerprint": "source", "conclusion": status, "basis": "已核对原文及计算依据", "source_refs": [{"quote": "原文摘录"}]}}


def test_三层报告统计一致且不适用不算问题(验证目录):
    results = [record("a", "确认错误"), record("b", "不适用"), record("c", "待核实")]
    data = {"meta": {"stage": "final", "company": "甲公司", "scope": "deep"}, "coverage": {"expected_ids": ["a", "b", "c"]}, "results": results}
    path = 验证目录 / "复核.xlsx"
    md = 验证目录 / "复核.md"
    export.write_excel(data, 验证目录 / "结果.json", path)
    export.write_markdown(data, md)
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["问题概览", "问题明细", "技术留痕"]
    assert wb["技术留痕"].sheet_state == "hidden"
    all_text = " ".join(str(c.value) for row in wb["问题概览"] for c in row)
    assert "致命" in all_text and "重要" in all_text
    text = md.read_text(encoding="utf-8-sig")
    assert "致命：1" in text and "重要：1" in text
    assert wb["技术留痕"].max_row == 4
    wb.close()


def test_未终审不得生成最终报告(验证目录):
    data = {"meta": {"stage": "final"}, "coverage": {"expected_ids": ["a"]}, "results": [{"check_id": "a", "passed": True, "severity": "info"}]}
    with pytest.raises(ValueError, match="未完成"):
        export.write_excel(data, 验证目录 / "结果.json", 验证目录 / "报告.xlsx")


def test_原文等号不作为Excel公式执行(验证目录):
    item = record("a", "确认错误")
    item["context"] = "=1+1"
    data = {"meta": {"stage": "final"}, "coverage": {"expected_ids": ["a"]}, "results": [item]}
    path = 验证目录 / "复核.xlsx"
    export.write_excel(data, 验证目录 / "结果.json", path)
    wb = load_workbook(path, data_only=False)
    assert all(c.data_type != "f" for ws in wb for row in ws for c in row)
    wb.close()


def test_位置使用日常名称且保留零号表索引():
    assert "货币资金" in export._location({"source_file": "报告.docx", "table_id": 0, "chapter": "货币资金"})
    assert "文档内第1个表格" in export._location({"source_file": "报告.docx", "table_id": 0})


def test_最终来源清单也保护原件(验证目录):
    source = 验证目录 / "原件.docx"
    source.write_bytes(b"original")
    data = {"meta": {"source_files": [{"path": str(source)}]}, "results": []}
    with pytest.raises(ValueError, match="覆盖"):
        export._protect(data, source)


def test_长中文依据行高能够显示(验证目录):
    item = record("a", "确认错误")
    item["expected"] = None
    item["actual"] = None
    item["difference"] = None
    item["review"]["basis"] = "第一处疑点：逐项核实金额和取数口径；第二处疑点：逐项核实金额和取数口径。"
    data = {"meta": {"stage": "final"}, "coverage": {"expected_ids": ["a"]}, "results": [item]}
    path = 验证目录 / "长内容.xlsx"
    export.write_excel(data, 验证目录 / "结果.json", path)
    wb = load_workbook(path)
    heights = [dim.height for dim in wb["问题明细"].row_dimensions.values() if dim.height]
    assert max(heights) > 48
    wb.close()


def test_跨表对象的表注差异进入问题明细(验证目录):
    item = record("a", "确认错误")
    item["check_type"] = "报表间"
    item["calculations"] = [{"check_type": "表注", "rule_name": "应收表注", "expected": "200", "actual": "199", "difference": "1", "calculation": {"match": False}}]
    data = {"meta": {"stage": "final"}, "coverage": {"expected_ids": ["a"]}, "results": [item]}
    path = 验证目录 / "分类.xlsx"
    export.write_excel(data, 验证目录 / "结果.json", path)
    wb = load_workbook(path)
    detail = "\n".join(str(c.value or "") for row in wb["问题明细"].iter_rows() for c in row)
    assert "应收表注" in detail
    wb.close()
