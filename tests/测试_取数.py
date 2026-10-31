import importlib
import json
import sys
from pathlib import Path
import pytest
from extractors.docx import find_table_in_extracted, iter_docx_rows
from extractors.base import extract_linear

def test_母公司来源精确匹配():
    data = {"tables": [{"id": 10, "source_file": "合并.docx", "rows": [["资产总计", "100"]]}, {"id": 11, "source_file": "母公司.docx", "rows": [["资产总计", "50"]]}]}
    assert find_table_in_extracted(data, {"source_file": "母公司.docx", "table_id": 11})["id"] == 11

def test_缺少精确定位不能猜第一张():
    data = {"tables": [{"id": 1, "name": "资产负债表", "rows": []}, {"id": 2, "name": "资产负债表", "rows": []}]}
    with pytest.raises(ValueError):
        find_table_in_extracted(data, "资产负债表")

def test_第一行实质数据不丢失():
    data = {"tables": [{"id": 0, "headers": ["货币资金", "100", "50"], "rows": [["资产总计", "100", "50"]]}]}
    rows = list(iter_docx_rows(data, [0]))
    assert rows[0][0] == "货币资金"

def test_同名非零冲突不得覆盖():
    with pytest.raises(ValueError, match="冲突"):
        extract_linear([["货币资金", "100", "50"], ["货币资金", "200", "50"]], 0, 1, 2)

def test_空合计不能自动作零():
    result = extract_linear([["资产总计", "", ""]], 0, 1, 2)
    assert not result.get("资产总计") or result["资产总计"].get("current") is None

def test_原始值不提前舍入():
    assert extract_linear([["金额", 1.23456, 0]], 0, 1, 2)["金额"]["current"] == "1.23456"

def test_Word正文页眉内容控件完整(验证目录, monkeypatch):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.add_paragraph("企业主体及会计期间应由当前agent判断")
    doc.sections[0].header.paragraphs[0].text = "甲公司页眉"
    sdt = OxmlElement("w:sdt"); content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p"); r = OxmlElement("w:r"); t = OxmlElement("w:t")
    t.text = "内容控件中的披露不能遗漏"; r.append(t); p.append(r); content.append(p); sdt.append(content)
    doc.element.body.insert(0, sdt)
    table = doc.add_table(rows=2, cols=3)
    for i,row in enumerate([["项目", "本期", "上期"], ["货币资金", "100", "50"]]):
        for j,value in enumerate(row): table.cell(i,j).text=value
    source = 验证目录 / "报告.docx"; doc.save(source)
    out = 验证目录 / "提取"
    import parse_report
    monkeypatch.setattr(sys,"argv",["parse_report.py",str(source),"-o",str(out)])
    parse_report.main()
    text=(out/"report.md").read_text(encoding="utf-8-sig")
    assert "甲公司页眉" in text and "内容控件中的披露不能遗漏" in text
    data=json.loads((out/"extracted_tables.json").read_text(encoding="utf-8-sig"))
    assert data["files"][0]["sha256"] and data["tables"][0]["rows"][0][0] == "项目"

def test_Excel公式缺缓存明确登记(验证目录, monkeypatch):
    from openpyxl import Workbook
    import parse_report
    wb=Workbook(); ws=wb.active; ws.title="报表"
    ws.append(["项目","本期","上期"]); ws.append(["金额","=1+2",3])
    source=验证目录/"报表.xlsx"; wb.save(source)
    out=验证目录/"提取"
    monkeypatch.setattr(sys,"argv",["parse_report.py",str(source),"-o",str(out)])
    parse_report.main()
    data=json.loads((out/"extracted_tables.json").read_text(encoding="utf-8-sig"))
    assert any("缓存" in x.get("reason","") for x in data["pending"])
    assert data["tables"][0]["rows"][1][1] is None

def test_隐藏工作表不作为复核资料(验证目录, monkeypatch):
    from openpyxl import Workbook
    import parse_report
    wb=Workbook(); visible=wb.active; visible.title="报出报表"
    visible.append(["项目","本期"]); visible.append(["经营活动现金流量净额","100"])
    hidden=wb.create_sheet("现补表"); hidden.sheet_state="hidden"
    hidden.append(["项目","本期"]); hidden.append(["经营活动现金流量净额","0"])
    source=验证目录/"报表.xlsx"; wb.save(source)
    out=验证目录/"提取"
    monkeypatch.setattr(sys,"argv",["parse_report.py",str(source),"-o",str(out)])
    parse_report.main()
    data=json.loads((out/"extracted_tables.json").read_text(encoding="utf-8-sig"))
    assert [t.get("sheet") for t in data["tables"]] == ["报出报表"]
    assert all(t.get("sheet") != "现补表" for t in data["tables"])
    assert any(x.get("sheet") == "现补表" and "隐藏" in x.get("reason","") for x in data["pending"])

def test_隐藏行与隐藏列不读取(验证目录, monkeypatch):
    from openpyxl import Workbook
    import parse_report
    wb=Workbook(); ws=wb.active; ws.title="现流表"
    ws.append(["项目","本期","内部列"])
    ws.append(["经营活动现金流量净额","100","999"])
    ws.append(["隐藏行内容","0","888"])
    ws.row_dimensions[3].hidden=True
    ws.column_dimensions["C"].hidden=True
    source=验证目录/"报表.xlsx"; wb.save(source)
    out=验证目录/"提取"
    monkeypatch.setattr(sys,"argv",["parse_report.py",str(source),"-o",str(out)])
    parse_report.main()
    data=json.loads((out/"extracted_tables.json").read_text(encoding="utf-8-sig"))
    text=json.dumps(data,ensure_ascii=False)
    assert "隐藏行内容" not in text
    assert "888" not in text
    assert "999" not in text
    assert "100" in text
    assert any("隐藏" in x.get("reason","") for x in data["pending"])

def test_图片只登记待识别不自动上传(验证目录, monkeypatch):
    import pymupdf as fitz, parse_report
    doc=fitz.open(); page=doc.new_page(); pix=page.get_pixmap()
    source=验证目录/"报表.png"; pix.save(source); doc.close()
    def deny(*args,**kwargs): raise AssertionError("不得自动执行外部OCR")
    monkeypatch.setattr(parse_report.subprocess,"run",deny)
    out=验证目录/"提取"
    monkeypatch.setattr(sys,"argv",["parse_report.py",str(source),"-o",str(out)])
    parse_report.main()
    data=json.loads((out/"extracted_tables.json").read_text(encoding="utf-8-sig"))
    assert data["pending"] and data["sources"] == ["报表.png"]

def test_单页扫描也不得漏记(验证目录, monkeypatch):
    import pymupdf as fitz, parse_report
    doc=fitz.open(); doc.new_page(); source=验证目录/"扫描.pdf"; doc.save(source); doc.close()
    monkeypatch.setattr(parse_report.subprocess,"run",lambda *a,**k: (_ for _ in ()).throw(AssertionError("不得上传")))
    out=验证目录/"提取"
    monkeypatch.setattr(sys,"argv",["parse_report.py",str(source),"-o",str(out)])
    parse_report.main()
    data=json.loads((out/"extracted_tables.json").read_text(encoding="utf-8-sig"))
    assert any(x.get("page")==1 for x in data["pending"])

def test_契约直接使用零号表并保留来源(验证目录):
    from 复核契约 import save_json
    import apply_manifest
    source=验证目录/"原件.docx";source.write_bytes(b"source")
    save_json(验证目录/"extracted_tables.json",{"source_root":str(验证目录),"tables":[{"id":0,"source_file":"原件.docx","row_model":"all_raw","headers":[],"rows":[["项目","本期","上期"],["货币资金","1.2345","0"]]}]})
    manifest={"schema_version":2,"files":{"原件.docx":{"loader":"docx"}},"statements_map":{"母公司资产负债表":{"source":"原件.docx","loader":"docx","source_table":0,"kind":"linear","first_data_row":1,"item_col":0,"current_col":1,"prior_col":2,"entity":"母公司","currency":"CNY","unit":"万元"}}}
    save_json(验证目录/"manifest.json",manifest)
    result=apply_manifest.apply_manifest(验证目录/"manifest.json")
    statement=result["母公司资产负债表"]
    assert statement["货币资金"]["current"]=="1.2345"
    assert statement["_source_map"]["货币资金"]["current"]["row"]==1
    assert statement["_meta"]["unit"]=="万元"

def test_续表冲突必须报错(验证目录):
    from 复核契约 import save_json
    import apply_manifest
    (验证目录/"原件.docx").write_bytes(b"source")
    save_json(验证目录/"extracted_tables.json",{"tables":[{"id":0,"source_file":"原件.docx","row_model":"all_raw","rows":[["货币资金","100","50"]]},{"id":1,"source_file":"原件.docx","row_model":"all_raw","rows":[["货币资金","200","50"]]}]})
    common={"source":"原件.docx","loader":"docx","kind":"linear","item_col":0,"current_col":1,"prior_col":2}
    save_json(验证目录/"manifest.json",{"files":{"原件.docx":{"loader":"docx"}},"statements_map":{"资产负债表":{**common,"source_table":0},"续表":{**common,"source_table":1,"merge_into":"资产负债表"}}})
    with pytest.raises(ValueError,match="冲突"):
        apply_manifest.apply_manifest(验证目录/"manifest.json")
