from openpyxl import load_workbook

import export_report as export
from 复核契约 import save_json


def issue(cid, status, basis, **extra):
    item = {
        "check_id": cid,
        "status": status,
        "rule_name": "table-" + cid,
        "check_type": "表注",
        "source_file": "D:/项目/03报出报表.xlsx",
        "table_id": 17,
        "row_range": [0, 8],
        "context": "科控集团3,500,000.00；科控酒店593,600.00",
        "expected": "4093600.00",
        "actual": "4060000.00",
        "difference": "33600.00",
        "calculation": {"match": status != "通过"},
        "review": {
            "agent": "当前agent",
            "check_id": cid,
            "source_fingerprint": "internal-fingerprint",
            "conclusion": status,
            "basis": basis,
            "source_refs": [{"quote": "科控酒店593,600.00"}],
        },
    }
    item.update(extra)
    return item


def visible_text(wb):
    return "\n".join(
        str(cell.value or "")
        for ws in wb.worksheets if ws.sheet_state == "visible"
        for row in ws.iter_rows() for cell in row
    )


def test_三层报告让人直接看懂且机器信息默认隐藏(验证目录):
    items = [
        issue("error1", "确认错误", "关联方应收明细比合计多33,600.00元，确认错误。", issue_title="关联方应收明细与合计不符", fix_suggestion="核实并更正关联方明细或应收账款合计。"),
        issue("pending1", "待核实", "封面编号仍为占位符，需项目组核实。"),
        issue("false1", "已排除误报", "差异系同笔重分类，已排除误报。"),
    ]
    data = {"meta": {"stage": "final", "company": "甲公司"}, "coverage": {"expected_ids": [x["check_id"] for x in items]}, "results": items}
    path = 验证目录 / "复核.xlsx"
    export.write_excel(data, 验证目录 / "results.json", path)
    wb = load_workbook(path, data_only=False)

    assert wb.sheetnames == ["问题概览", "问题明细", "技术留痕"]
    assert wb["技术留痕"].sheet_state == "hidden"
    text = visible_text(wb)
    for required in ("致命", "重要", "细节", "关联方应收明细与合计不符", "错在哪", "原文", "核对结果", "凭什么说错", "怎么改", "03报出报表.xlsx", "4,093,600.00", "4,060,000.00", "33,600.00"):
        assert required in text
    for machine in ("error1", "pending1", "internal-fingerprint", "D:/项目/", "results.json", "table-error1", "row_range", "表格ID", "对象ID"):
        assert machine not in text
    assert wb["技术留痕"]["A2"].value == "error1"
    assert wb["问题明细"].max_row == 15
    wb.close()


def test_位置使用真实工作表名并展示对应原始行(验证目录):
    folder = 验证目录 / "输入"
    folder.mkdir()
    save_json(folder / "extracted_tables.json", {
        "files": [],
        "tables": [{"id": 8, "source_file": "03报出报表.xlsx", "sheet": "现金流量表补充资料", "rows": [["项目", "本期金额"], ["经营活动现金流量净额", "0"], ["主表经营净额", "11449812.43"]]}],
    })
    item = issue("error1", "确认错误", "补充资料未填列，确认错误。", table_id=8, row_range=[0, 2], context='"row_range": [0, 2]')
    item["calculations"] = [{"rule_name": "补充资料经营净额与主表不符", "expected": "11449812.43", "actual": "0", "difference": "11449812.43", "operands": [{"table_id": 8, "row": 2, "col": 1}], "target": {"table_id": 8, "row": 1, "col": 1}, "calculation": {"match": False}}]
    data = {"meta": {"stage": "final", "input_dir": str(folder)}, "coverage": {"expected_ids": ["error1"]}, "results": [item]}
    path = 验证目录 / "复核.xlsx"
    export.write_excel(data, 验证目录 / "results.json", path)
    wb = load_workbook(path)
    text = visible_text(wb)
    assert "现金流量表补充资料" in text
    assert "经营活动现金流量净额" in text
    assert "主表经营净额" in text
    assert "文档内第9个表格" not in text
    wb.close()


def test_相同未执行原因归并且不淹没主要问题(验证目录):
    items = [issue("u1", "未执行", "缺少必要资料，尚未完成核对"), issue("u2", "未执行", "缺少必要资料，尚未完成核对", table_id=18)]
    data = {"meta": {"stage": "draft"}, "coverage": {"expected_ids": ["u1", "u2"]}, "results": items}
    path = 验证目录 / "复核.xlsx"
    export.write_excel(data, 验证目录 / "results.json", path)
    wb = load_workbook(path)
    ws = wb["问题明细"]
    assert sum(cell.value == "问题" for row in ws.iter_rows() for cell in row if cell.column == 1) == 1
    assert "共2条同类记录" in visible_text(wb)
    wb.close()


def test_Markdown只保留人类阅读内容(验证目录):
    item = issue("error1", "确认错误", "关联方应收明细比合计多33,600.00元，确认错误。", issue_title="关联方应收明细与合计不符", fix_suggestion="核实并更正关联方明细或应收账款合计。")
    data = {"meta": {"stage": "final", "company": "甲公司"}, "coverage": {"expected_ids": ["error1"]}, "results": [item]}
    path = 验证目录 / "复核.md"
    export.write_markdown(data, path)
    text = path.read_text(encoding="utf-8-sig")
    for required in ("问题概览", "问题明细", "错在哪", "原文", "凭什么说错", "怎么改", "关联方应收明细与合计不符"):
        assert required in text
    for machine in ("error1", "internal-fingerprint", "D:/项目/", "table-error1", "对象ID", "技术留痕"):
        assert machine not in text


def test_人工严重程度优先于程序状态(验证目录):
    item = issue("minor1", "确认错误", "附注格式存在一处细节问题。")
    item["review"]["severity"] = "细节"
    data = {"meta": {"stage": "final"}, "coverage": {"expected_ids": ["minor1"]}, "results": [item]}
    path = 验证目录 / "复核.xlsx"
    export.write_excel(data, 验证目录 / "结果.json", path)
    wb = load_workbook(path)
    assert "细节｜" in visible_text(wb)
    assert "核对关系：应为" in visible_text(wb)
    assert "致命｜" not in visible_text(wb)
    wb.close()


def test_同源错误合并且内部关系编号不残留(验证目录):
    first = issue("error1", "确认错误", "unaud-cfs-op-cur/unaud-cfs-fin-cur属已排除误报；supp-op-net-cur/supp-op-net-pri补充资料空白，确认错误。")
    first["calculations"] = [
        {"rule_name": "未审CFS经营净额本期vs报出(预期存在重分类调整差异)", "expected": "36249812.43", "actual": "11449812.43", "difference": "24800000", "calculation": {"match": False}},
        {"rule_name": "现补表经营净额本期=CFS主表本期(预期不符)", "expected": "11449812.43", "actual": "0", "difference": "11449812.43", "calculation": {"match": False}},
        {"rule_name": "现补表经营净额上期=CFS主表上期(预期不符)", "expected": "-19487512.21", "actual": "0", "difference": "19487512.21", "calculation": {"match": False}},
    ]
    second = issue("error2", "确认错误", "补充资料空白，确认错误。", calculations=first["calculations"][1:])
    summary = issue("summary", "确认错误", "跨表关系汇总", use_declared_relations=True)
    routine = issue("routine", "未执行", "明细中段，无小计行，本段无独立验算对象")
    items = [first, second, summary, routine]
    data = {"meta": {"stage": "draft"}, "coverage": {"expected_ids": [x["check_id"] for x in items]}, "results": items}
    path = 验证目录 / "复核.xlsx"
    export.write_excel(data, 验证目录 / "结果.json", path)
    wb = load_workbook(path)
    text = visible_text(wb)
    assert sum(cell.value == "问题" for row in wb["问题明细"].iter_rows() for cell in row if cell.column == 1) == 1
    assert "现补表经营净额" in text
    for residue in ("unaud", "supp-op", "rpart-", "/-", "cross_relations", "明细中段，无小计行", "：/", "；/"):
        assert residue not in text
    assert "未审CFS经营净额" not in text
    wb.close()
