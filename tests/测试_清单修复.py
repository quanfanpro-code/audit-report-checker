"""16项审计发现的行为回归；全部使用新建合成资料。"""
from pathlib import Path
import copy
import pytest
import 复核任务 as task
import 复核契约 as contract
import run_check as run
import parse_report as parse
import export_report as export
import calculator as calc
import calculator_rules as rules

def prepare(folder,plan="na"):
    folder.mkdir();source=folder/"原件.txt";contract.save_text(source,"合成原件，不得覆盖")
    tables=[{"id":i,"source_file":source.name,"headers":[],"row_model":"all_raw","chapter":name,"unit":"元","entity":"单体","currency":"CNY","rows":[["项目","本期金额"],["现金","100"],["应收账款",str(ar)],["合计",str(100+ar)]]} for i,name,ar in ((0,"报表",201),(1,"附注",200))]
    contract.save_json(folder/"extracted_tables.json",{"tables":tables,"sources":[source.name],"files":[{"name":source.name,"path":str(source),"sha256":contract.fingerprint(source)}],"segments":[],"pending":[]})
    contract.save_text(folder/"report.md","合成报告")
    if plan is not None:contract.save_json(folder/"relations.json",{"checks":[],"not_applicable_basis":"本测试只验证资料队列，不作跨表业务结论"} if plan=="na" else plan)
    return source

def reply(state,batch,status="不适用"):
    objects={o["check_id"]:o for o in state["objects"]}
    return {"batch_id":batch["batch_id"],"source_fingerprint":batch["source_fingerprint"],"results":[{"check_id":cid,"status":status,"review":{"agent":"当前agent","check_id":cid,"source_fingerprint":batch["source_fingerprint"],"conclusion":status,"basis":"合成回归核对","source_refs":[{"quote":objects[cid]["material"][:20]}]}} for cid in batch["check_ids"]]}

def relation(name,row):
    meta={"entity":"单体","period":"current","currency":"CNY","unit":"元"}
    return {"check_id":name,"rule_name":name,"entity":"单体","period":"current","currency":"CNY","formula":"x","variables":{"x":{"table_id":0,"row":row,"col":1,**meta}},"target":{"table_id":1,"row":row,"col":1,**meta}}

@pytest.mark.parametrize("target",["原件.txt","report.md","extracted_tables.json","relations.json","复核进度.json"])
def test_01_下一批不能覆盖来源和关键输入(验证目录,target):
    folder=验证目录/"输入";prepare(folder);task.init_review(folder,"128k","deep")
    path=folder/target;before=path.read_bytes()
    with pytest.raises(SystemExit):task.main(["下一批",str(folder),"-o",str(path)])
    assert path.read_bytes()==before

def test_02_缺少应检关系不允许用一条计算盖章(验证目录):
    folder=验证目录/"输入";checks=[relation("现金",1),relation("应收账款",2)];prepare(folder,{"checks":checks})
    state=task.init_review(folder,"128k","deep");batch=task.next_batch(folder);result=reply(state,batch)
    cid=next(o["check_id"] for o in state["objects"] if o["kind"]=="cross_relations")
    row=next(r for r in result["results"] if r["check_id"]==cid)
    row.update(status="通过",relations=checks[:1]);row["review"]["conclusion"]="通过"
    with pytest.raises(ValueError,match="应检|范围|关系清单"):task.submit_batch(folder,result)

def test_02_没有关系清单不能宣称跨表通过(验证目录):
    folder=验证目录/"输入";prepare(folder,None);state=task.init_review(folder,"128k","deep");batch=task.next_batch(folder);result=reply(state,batch)
    cid=next(o["check_id"] for o in state["objects"] if o["kind"]=="cross_relations")
    row=next(r for r in result["results"] if r["check_id"]==cid);row.update(status="通过",relations=[relation("现金",1)]);row["review"]["conclusion"]="通过"
    with pytest.raises(ValueError,match="清单|声明|未执行"):task.submit_batch(folder,result)

def test_02_完整清单执行后保留真实差异(验证目录):
    folder=验证目录/"输入";checks=[relation("现金",1),relation("应收账款",2)];prepare(folder,{"checks":checks})
    state=task.init_review(folder,"128k","deep");batch=task.next_batch(folder);result=reply(state,batch)
    cid=next(o["check_id"] for o in state["objects"] if o["kind"]=="cross_relations")
    row=next(r for r in result["results"] if r["check_id"]==cid)
    row.update(status="确认错误",use_declared_relations=True);row["review"]["conclusion"]="确认错误"
    task.submit_batch(folder,result);final=task.finalize_results(folder,folder/"结果.json")
    assert final["summary"]["完整"] and final["summary"]["确认错误"]==1
    item=next(r for r in final["results"] if r["check_id"]==cid)
    assert len(item["calculations"])==2 and str(item["calculations"][1]["difference"])=="1"

def test_04_原生完整声明可以正常验算():
    table={"id":0,"row_model":"all_raw","unit":"元","rows":[["项目","本期余额","本期准备","本期净值"],["甲","100","20","80"]]}
    ann={"header_rows":[0],"columns":[{"col":c,"type":"amount","period":"current"} for c in (1,2,3)],"horizontal":[{"rows":[1],"operands":[{"col":1,"op":"+"},{"col":2,"op":"-"}],"result_col":3}]}
    assert run.execute_annotations(table,ann)[0]["calculation"]["match"] is True

def test_03_引用不能覆盖来源单位():
    tables={0:{"id":0,"unit":"万元","entity":"单体","period":"current","currency":"CNY","rows":[["现金","100"]]},1:{"id":1,"unit":"元","entity":"单体","period":"current","currency":"CNY","rows":[["现金","100"]]}}
    rule=relation("单位",0)
    with pytest.raises(ValueError,match="单位"):run.execute_relation(rule,tables)

@pytest.mark.parametrize("headers,values",[(["项目","本期期末余额","本期准备","上期期末净值"],["甲","100","20","80"]),(["项目","金额","计提比例","合计"],["甲","100","0.1","100.1"])])
def test_04_原生表格必须有列角色且不能跨期或加比例(headers,values):
    table={"id":0,"row_model":"all_raw","headers":[],"unit":"元","rows":[headers,values]}
    ann={"horizontal":[{"rows":[1],"operands":[{"col":1,"op":"+"},{"col":2,"op":"-"}],"result_col":3}]}
    with pytest.raises(ValueError,match="列|表头|期间|比例"):run.execute_annotations(table,ann)
    ann.update(header_rows=[0],columns=[{"col":c,"type":"amount","period":"current"} for c in (1,2,3)])
    with pytest.raises(ValueError,match="期间|时点|比例|百分比"):run.execute_annotations(table,ann)

def test_05_未执行后可重新提交且旧提交仍幂等(验证目录):
    folder=验证目录/"输入";prepare(folder);state=task.init_review(folder,"128k","deep");batch=task.next_batch(folder)
    first=reply(state,batch,"未执行");task.submit_batch(folder,first)
    state=task.init_review(folder,"128k","deep");again=task.next_batch(folder);new=reply(state,again)
    task.submit_batch(folder,new)
    assert task.audit_coverage(folder)["未执行"]==0
    assert task.submit_batch(folder,new)==task.submit_batch(folder,new)

def test_06_原件变更后导出拒绝旧结论(验证目录):
    source=验证目录/"源件.txt";contract.save_text(source,"原版");files=[{"path":str(source),"sha256":contract.fingerprint(source)}]
    data={"meta":{"stage":"draft","source_files":files},"coverage":{"expected_ids":[]},"results":[]}
    contract.save_json(验证目录/"结果.json",data);contract.save_text(source,"新版")
    for fn in (lambda:export.load_results(验证目录/"结果.json"),lambda:export.write_markdown(data,验证目录/"报告.md"),lambda:export.write_excel(data,验证目录/"结果.json",验证目录/"报告.xlsx")):
        with pytest.raises(ValueError,match="来源|原件|变化|失效"):fn()

def test_07_权益禁止自比较并报告缺期():
    matrix={"rows":[["期末","100","50","150"]],"columns":["股本","未分配利润","合计"],"beginning_row":0,"ending_row":0,"change_rows":[],"total_col":3}
    rows=rules.check_equity_change_statement({"_matrix":{"current":matrix}})
    assert rows and all(not r.passed for r in rows)
    assert any("prior" in r.rule_name or "上期" in r.rule_name for r in rows)

def test_08_商誉规则区分吸收合并(技能根目录):
    text=(技能根目录/"references/rules.md").read_text(encoding="utf-8-sig")
    assert "合并报表才有商誉" not in text and "吸收合并" in text

def test_09_交易性与衍生资产不机械相加(技能根目录):
    text=(技能根目录/"references/rules.md").read_text(encoding="utf-8-sig")
    assert "衍生金融资产 + 其他交易性金融资产" not in text
    assert "分别列示" in text or "分别映射" in text

def test_10_Word制表符换行和表内段落不丢(验证目录):
    from docx import Document
    doc=Document();doc.add_paragraph("收入\t100\n成本\t80");t=doc.add_table(rows=1,cols=1);t.cell(0,0).text="甲\n乙";source=验证目录/"原件.docx";doc.save(source)
    parse.main([str(source),"-o",str(验证目录/"提取")]);data=contract.load_json(验证目录/"提取/extracted_tables.json")
    assert data["segments"][0]["text"]=="收入\t100\n成本\t80"
    assert data["tables"][0]["rows"][0][0]=="甲\n乙"

@pytest.mark.parametrize("value",["1 2","1\n2","1\t2"])
def test_11_歧义金额不能直接拼接(value):
    with pytest.raises(ValueError):calc.parse_number(value)

@pytest.mark.parametrize("value",["1,2","1,,2"])
def test_11_表达式不能删逗号改写(value):
    with pytest.raises(ValueError):calc.evaluate_expression(value)

def test_12_OCR失败仍保留本地正文(验证目录,monkeypatch):
    from docx import Document
    doc=Document();doc.add_paragraph("本地正文不能丢");source=验证目录/"原件.docx";doc.save(source)
    monkeypatch.setattr(parse.shutil,"which",lambda _:None)
    parse.main([str(source),"-o",str(验证目录/"提取"),"--allow-cloud-ocr"])
    data=contract.load_json(验证目录/"提取/extracted_tables.json")
    assert any("本地正文不能丢" in s["text"] for s in data["segments"])
    assert data["pending"] and data["files"][0]["status"]!="提取失败"

def test_13_大表每片都有完整行和背景(验证目录):
    folder=验证目录/"输入";prepare(folder);data=contract.load_json(folder/"extracted_tables.json")
    table=data["tables"][0];table.update(chapter="合并应收明细",header_rows=[0]);table["rows"]=[["项目","本期金额"]]+[["项目"+str(i),"100"] for i in range(3000)]
    contract.save_json(folder/"extracted_tables.json",data);state=task.init_review(folder,"128k","deep")
    parts=[o for o in state["objects"] if o.get("table_id")==0]
    assert len(parts)>1;seen=[]
    for part in parts:
        assert "本期金额" in part["material"] and "合并应收明细" in part["material"] and "单体" in part["material"]
        start,end=part["row_range"];seen.extend(range(start,end))
        assert len(part["material"])<6000
    assert sorted(set(seen))==list(range(len(table["rows"])))

def test_14_旧Excel纳入来源与待办(验证目录):
    from docx import Document
    folder=验证目录/"原始资料";folder.mkdir();doc=Document();doc.add_paragraph("报告");doc.save(folder/"报告.docx");contract.save_text(folder/"财务报表.xls","不支持格式的合成样本")
    parse.main([str(folder),"-o",str(验证目录/"提取")]);data=contract.load_json(验证目录/"提取/extracted_tables.json")
    assert "财务报表.xls" in data["sources"] and any(x.get("source_file")=="财务报表.xls" for x in data["pending"])

def test_13_已声明口径随分片传递(验证目录):
    folder=验证目录/"输入";prepare(folder,{"checks":[relation("现金",1)]})
    data=contract.load_json(folder/"extracted_tables.json")
    for table in data["tables"]:
        for key in ("unit","currency","entity"):table.pop(key)
    contract.save_json(folder/"extracted_tables.json",data)
    state=task.init_review(folder,"128k","deep")
    assert all('"entity": "单体"' in o["material"] and '"unit": "元"' in o["material"] for o in state["objects"] if o["kind"]=="table")

def test_15_写入中断保留旧进度(验证目录,monkeypatch):
    path=验证目录/"进度.json";contract.save_json(path,{"旧进度":1});old=path.read_bytes();original=Path.open
    class Broken:
        def __init__(self,f):self.f=f
        def __enter__(self):return self
        def __exit__(self,*args):self.f.close()
        def write(self,text):self.f.write(text[:5]);self.f.flush();raise OSError("模拟中断")
        def flush(self):self.f.flush()
        def fileno(self):return self.f.fileno()
    def open_file(p,mode="r",*args,**kwargs):
        f=original(p,mode,*args,**kwargs)
        return Broken(f) if p.parent==path.parent and any(x in mode for x in ("w","x")) else f
    monkeypatch.setattr(Path,"open",open_file)
    with pytest.raises(OSError):contract.save_json(path,{"新进度":2})
    assert path.read_bytes()==old

def test_16_运行目录不再保留旧模型入口和死函数(技能根目录):
    import ast
    assert not (技能根目录/"scripts/ai_worker.py").exists()
    source=(技能根目录/"scripts/run_check.py").read_text(encoding="utf-8-sig")
    assert all(flag not in source for flag in ("--no-ai","--api-key","--model","auto-manifest","skip-text-ai"))
    names=set()
    for path in (技能根目录/"scripts").rglob("*.py"):
        names.update(n.name for n in ast.walk(ast.parse(path.read_text(encoding="utf-8-sig"))) if isinstance(n,ast.FunctionDef))
    assert not names.intersection({"clean_number_cell","extract_matrix","iter_xlsx_rows","effective_severity"})

def test_01_取数输出不能覆盖进度(验证目录):
    import apply_manifest
    folder=验证目录/"输入";prepare(folder);task.init_review(folder,"128k","deep")
    contract.save_json(folder/"manifest.json",{"files":{"原件.txt":{"loader":"docx"}},"statements_map":{}})
    target=folder/"复核进度.json";before=target.read_bytes()
    with pytest.raises(SystemExit):apply_manifest.main([str(folder/"manifest.json"),"-o",str(target)])
    assert target.read_bytes()==before

def test_01_导出不能覆盖提取正文(验证目录):
    folder=验证目录/"输入";prepare(folder);before=(folder/"report.md").read_bytes()
    data={"meta":{"stage":"draft","input_dir":str(folder)},"coverage":{"expected_ids":[]},"results":[]}
    with pytest.raises(ValueError,match="覆盖"):export.write_markdown(data,folder/"report.md")
    assert (folder/"report.md").read_bytes()==before
